"""Carga de las tres fuentes de reservas y normalización de fechas.

Fuentes (y sus trampas):
  * Reservas 2023-2024 (xlsx): una hoja; Día/Año numéricos; fechas de estancia; 8 filas de 2024 al final.
  * Reservas 2024 Actualizado (xlsx): 12 hojas mensuales + "Consolidado 2024". El Consolidado NO es
    confiable: las 14 filas de octubre quedan corridas una columna ("Octubre" en Día, "2024" en Mes) y sin
    "Mes del Reporte". Por eso se leen las hojas mensuales y el Consolidado solo se usa para auditar.
    El nombre de la hoja es el mes en que se reportó la reserva; la estancia es Día/Mes/Año.
  * 2026 RESERVAS (pdf): tabla Nombre | Teléfono | mes. No trae día de estancia, solo el mes.
"""
from __future__ import annotations

import re
import unicodedata
from collections import Counter
from dataclasses import dataclass
from datetime import date
from pathlib import Path

import openpyxl
import pdfplumber

FUENTE_2023_2024 = "Reservas 2023-2024"
FUENTE_2024 = "Reservas 2024"
FUENTE_2026 = "Reservas 2026"

PRECISION_DIA = "día"
PRECISION_MES = "mes"
PRECISION_REPORTE = "solo mes de reporte"
PRECISION_NINGUNA = "sin fecha"

MESES = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6, "julio": 7,
    "agosto": 8, "septiembre": 9, "setiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12,
}


@dataclass
class Registro:
    """Una fila del origen, ya con la fecha interpretada (nombre/teléfono se normalizan después)."""

    fuente: str
    ubicacion: str  # para poder volver a la fila original: "hoja 'Marzo 2024', fila 5"
    nombre_raw: str
    telefono_raw: str
    fecha: date | None  # fecha exacta de estancia
    precision: str  # PRECISION_*
    anio: int | None  # año de la estancia (o del reporte si no hay otra cosa)
    mes: int | None  # mes de la estancia (o del reporte)
    periodo_reporte: str  # "2024-03" si se conoce el mes en que se reportó, si no ""
    nota_fecha: str = ""  # "No disponible", "Fecha abierta", "fecha inválida: 31/04/2024"...


def _sin_acentos(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def parse_mes(valor: object) -> int | None:
    """'de Junio', 'De Junio', 'enero', 'Septiembre' -> 6, 6, 1, 9. Cualquier otra cosa -> None."""
    s = _sin_acentos("" if valor is None else str(valor)).lower().strip()
    s = re.sub(r"^de\s+", "", s)
    return MESES.get(s)


def _exigir_columnas(cab: list[str], requeridas: list[str], donde: str) -> None:
    faltan = [c for c in requeridas if c not in cab]
    if faltan:
        raise ValueError(f"{donde}: faltan las columnas {faltan}; encabezados encontrados: {cab}")


def _a_int(valor: object) -> int | None:
    s = "" if valor is None else str(valor).strip()
    return int(s) if re.fullmatch(r"\d{1,4}", s) else None


def interpretar_fecha(dia: object, mes: object, anio: object) -> tuple[date | None, str, int | None, int | None, str]:
    """Devuelve (fecha, precisión, año, mes, nota). Nunca lanza excepción por un dato sucio."""
    d, m, a = _a_int(dia), parse_mes(mes), _a_int(anio)
    if m and a:
        if d:
            try:
                return date(a, m, d), PRECISION_DIA, a, m, ""
            except ValueError:
                return None, PRECISION_MES, a, m, f"fecha inválida: {d:02d}/{m:02d}/{a}"
        return None, PRECISION_MES, a, m, ""
    textos = {str(x).strip() for x in (dia, mes, anio) if x is not None and str(x).strip()}
    nota = next((t for t in textos if not t.isdigit() and parse_mes(t) is None), "")
    return None, PRECISION_NINGUNA, a, m, nota or "sin fecha"


# ------------------------------------------------------------------------------------------------
def cargar_2023_2024(path: str | Path) -> list[Registro]:
    ws = openpyxl.load_workbook(path, data_only=True).active
    filas = ws.iter_rows(values_only=True)
    cab = [str(c).strip() if c else "" for c in next(filas)]
    _exigir_columnas(cab, ["Nombre de la Reserva", "Teléfono", "Día", "Mes", "Año"], "Reservas 2023-2024")
    idx = {nombre: i for i, nombre in enumerate(cab)}
    out: list[Registro] = []
    for n, fila in enumerate(filas, start=2):
        if not fila or not (fila[idx["Teléfono"]] or fila[0]):
            continue
        fecha, prec, anio, mes, nota = interpretar_fecha(fila[idx["Día"]], fila[idx["Mes"]], fila[idx["Año"]])
        out.append(Registro(FUENTE_2023_2024, f"fila {n}", str(fila[0] or "").strip(),
                            str(fila[idx["Teléfono"]] or "").strip(), fecha, prec, anio, mes, "", nota))
    return out


def _periodo_de_hoja(titulo: str) -> tuple[int, int] | None:
    m = re.match(r"^\s*([A-Za-zÁÉÍÓÚáéíóúñÑ]+)\s+(\d{4})\s*$", titulo)
    if m and parse_mes(m.group(1)):
        return int(m.group(2)), parse_mes(m.group(1))  # type: ignore[return-value]
    return None


def cargar_2024(path: str | Path) -> tuple[list[Registro], dict[str, int]]:
    """Lee las hojas mensuales. Devuelve (registros, auditoría del Consolidado)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    out: list[Registro] = []
    for ws in wb.worksheets:
        per = _periodo_de_hoja(ws.title)
        if per is None:  # "Consolidado 2024" y cualquier hoja que no sea mensual
            continue
        filas = ws.iter_rows(values_only=True)
        cab = [str(c).strip() if c else "" for c in next(filas)]
        _exigir_columnas(cab, ["Teléfono", "Mes", "Año"], f"hoja '{ws.title}'")
        col_nombre = next((i for i, c in enumerate(cab) if c.startswith("Nombre")), None)
        if col_nombre is None:
            raise ValueError(f"hoja '{ws.title}': no hay columna de nombre; encabezados: {cab}")
        col_tel, col_mes, col_anio = cab.index("Teléfono"), cab.index("Mes"), cab.index("Año")
        col_dia = cab.index("Día") if "Día" in cab else None
        for n, fila in enumerate(filas, start=2):
            if not fila or not (fila[col_nombre] or fila[col_tel]):
                continue
            fecha, prec, anio, mes, nota = interpretar_fecha(
                fila[col_dia] if col_dia is not None else None, fila[col_mes], fila[col_anio])
            out.append(Registro(FUENTE_2024, f"hoja '{ws.title}', fila {n}", str(fila[col_nombre] or "").strip(),
                                str(fila[col_tel] or "").strip(), fecha, prec, anio, mes,
                                f"{per[0]}-{per[1]:02d}", nota))
    return out, auditar_consolidado(wb)


def auditar_consolidado(wb: "openpyxl.Workbook") -> dict[str, int]:
    """Cuenta filas del Consolidado cuyo Mes/Año no son un mes/año reales (columnas corridas)."""
    if "Consolidado 2024" not in wb.sheetnames:
        return {}
    filas = list(wb["Consolidado 2024"].iter_rows(min_row=2, values_only=True))
    corridas = sum(1 for f in filas if parse_mes(f[3]) is None and str(f[3]).strip().isdigit())
    sin_reporte = sum(1 for f in filas if not f[5])
    return {"consolidado_filas": len(filas), "consolidado_filas_corridas": corridas,
            "consolidado_sin_mes_reporte": sin_reporte}


def _celda(valor: object) -> str:
    return " ".join(str(valor or "").split())


def filas_de_pdf(path: str | Path) -> list[tuple[int, list[str]]]:
    """(página, [nombre, teléfono, mes]) de cada fila de la tabla, sin cabecera."""
    filas: list[tuple[int, list[str]]] = []
    with pdfplumber.open(path) as pdf:
        for num, pagina in enumerate(pdf.pages, start=1):
            for tabla in pagina.extract_tables():
                for fila in tabla:
                    celdas = [_celda(c) for c in fila]
                    if len(celdas) < 3 or not any(celdas) or celdas[0].lower() == "nombre":
                        continue
                    # El mes parte la palabra en dos líneas ("Septiembr\ne"): se une sin espacios.
                    celdas[2] = re.sub(r"\s+", "", str(fila[2] or ""))
                    filas.append((num, celdas))
    if not filas:
        raise ValueError(f"{path}: no se detectó ninguna tabla con filas. ¿El PDF es una imagen escaneada o "
                         "cambió su formato? Exporta de nuevo la tabla con líneas de cuadrícula.")
    return filas


def cargar_2026(path: str | Path, anio: int = 2026) -> list[Registro]:
    out: list[Registro] = []
    por_pagina: Counter[int] = Counter()
    for pagina, (nombre, telefono, mes_txt) in filas_de_pdf(path):
        por_pagina[pagina] += 1
        mes = parse_mes(mes_txt)
        out.append(Registro(FUENTE_2026, f"pág. {pagina}, fila {por_pagina[pagina]}", nombre, telefono, None,
                            PRECISION_REPORTE if mes else PRECISION_NINGUNA, anio if mes else None, mes,
                            f"{anio}-{mes:02d}" if mes else "", "" if mes else f"mes no reconocido: {mes_txt!r}"))
    return out
