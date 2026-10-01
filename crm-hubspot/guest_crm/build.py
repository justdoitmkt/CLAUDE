"""Une las tres fuentes en una base de huéspedes lista para importar a HubSpot.

Uso:
    python -m guest_crm.build \\
        --xlsx-2023-2024 Reservas_2023_2024.xlsx \\
        --xlsx-2024 Reservas_2024_Actualizado.xlsx \\
        --pdf-2026 2026_RESERVAS.pdf \\
        --out out/

Salidas (en --out):
    contactos_hubspot.csv     un renglón por teléfono único: es el archivo que se importa a HubSpot
    reservas_historial.csv    un renglón por reserva distinta (archivo de respaldo / para análisis)
    revision_manual.xlsx      lo que requiere una persona (sin teléfono, nombres dudosos, posibles duplicados)
    reporte_calidad.md        conteos y hallazgos (sin datos personales)

Todo contiene datos personales: `out/` está en .gitignore y no debe subirse a un repositorio.
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import hubspot_schema as H
from . import sources as S
from .names import ALTA, BAJA, INVALIDA, MEDIA, NombreLimpio, clave_comparacion, limpiar_nombre
from .phones import ASUMIDO_MX, REVISAR, SIN_TELEFONO, VALIDO, PhoneResult, normalize_phone

ORDEN_FUENTES = [S.FUENTE_2023_2024, S.FUENTE_2024, S.FUENTE_2026]
COLUMNAS_HISTORIAL = ["telefono_e164", "nombre", "fecha_estancia", "precision_fecha", "anio", "mes",
                      "periodos_reporte", "fuentes", "unidades", "filas_origen", "ubicaciones"]
_RANGO_CONFIANZA = {ALTA: 3, MEDIA: 2, BAJA: 1, INVALIDA: 0}
_ETIQUETA_CONFIANZA = {ALTA: "Alta", MEDIA: "Media", BAJA: "Baja", INVALIDA: "Sin nombre"}


# ------------------------------------------------------------------------------------------------
# Modelo
# ------------------------------------------------------------------------------------------------
@dataclass
class Fila:
    reg: S.Registro
    tel: PhoneResult
    nom: NombreLimpio


@dataclass
class Reserva:
    contacto_id: str
    fecha: date | None
    precision: str
    anio: int | None
    mes: int | None
    filas: list[Fila] = field(default_factory=list)

    @property
    def fuentes(self) -> list[str]:
        return [f for f in ORDEN_FUENTES if any(x.reg.fuente == f for x in self.filas)]

    @property
    def unidades(self) -> int:
        """Filas de la misma hoja para la misma fecha ("Segunda reserva") = más de una unidad ese día."""
        return max(Counter((x.reg.fuente, x.reg.periodo_reporte) for x in self.filas).values())

    @property
    def fecha_ref(self) -> date | None:
        if self.fecha:
            return self.fecha
        if self.anio and self.mes and self.precision in (S.PRECISION_MES, S.PRECISION_REPORTE):
            return date(self.anio, self.mes, 1)
        return None

    @property
    def anio_ref(self) -> int | None:
        if self.anio:
            return self.anio
        for x in self.filas:
            if x.reg.periodo_reporte:
                return int(x.reg.periodo_reporte[:4])
        return None

    def etiqueta(self) -> str:
        if self.precision == S.PRECISION_DIA and self.fecha:
            txt = self.fecha.isoformat()
        elif self.precision == S.PRECISION_MES and self.anio and self.mes:
            txt = f"{self.anio}-{self.mes:02d} (mes)"
        elif self.precision == S.PRECISION_REPORTE and self.anio and self.mes:
            txt = f"{self.anio}-{self.mes:02d} (mes de reporte)"
        else:
            nota = next((x.reg.nota_fecha for x in self.filas if x.reg.nota_fecha), "")
            txt = f"sin fecha ({nota})" if nota else "sin fecha"
        return txt + (f" ×{self.unidades}" if self.unidades > 1 else "")


@dataclass
class Contacto:
    id: str  # E.164, o "NOTEL::<nombre>" si no hay teléfono
    e164: str | None
    reservas: list[Reserva]
    filas: list[Fila]
    nombre: NombreLimpio
    confianza: str  # ALTA | MEDIA | BAJA | INVALIDA (ya considerando conflictos)
    variantes: list[str]
    conflicto: bool
    calidad: str
    notas: list[str]


def _clave_estancia(r: S.Registro) -> tuple:
    if r.fecha:
        return ("d", r.fecha.isoformat())
    if r.precision == S.PRECISION_MES:
        return ("m", f"{r.anio}-{r.mes:02d}")
    if r.precision == S.PRECISION_REPORTE:
        return ("r", r.periodo_reporte)
    return ("?", r.fuente, r.periodo_reporte, r.nota_fecha)


_PALABRAS_GENERICAS = {"familia", "fam", "boda", "empresa", "hijo", "dra", "dr", "lic", "ing", "segunda", "reserva"}


def _tokens(nombre: str) -> set[str]:
    return {t for t in clave_comparacion(nombre).split() if len(t) > 1 and t not in _PALABRAS_GENERICAS}


def _digitos_distintos(a: str, b: str) -> int | None:
    """Cuántos dígitos difieren entre dos teléfonos de igual longitud (1 = casi seguro error de captura)."""
    return sum(x != y for x, y in zip(a, b)) if len(a) == len(b) else None


# ------------------------------------------------------------------------------------------------
# Procesamiento
# ------------------------------------------------------------------------------------------------
def normalizar(registros: list[S.Registro]) -> list[Fila]:
    return [Fila(r, normalize_phone(r.telefono_raw), limpiar_nombre(r.nombre_raw)) for r in registros]


def agrupar_reservas(filas: list[Fila]) -> list[Reserva]:
    """Una reserva = mismo contacto + misma estancia, aunque aparezca en varias hojas o archivos."""
    reservas: dict[tuple, Reserva] = {}
    for f in filas:
        cid = f.tel.e164 or f"NOTEL::{clave_comparacion(f.nom.completo or f.reg.nombre_raw)}"
        clave = (cid, _clave_estancia(f.reg))
        if clave not in reservas:
            reservas[clave] = Reserva(cid, f.reg.fecha, f.reg.precision, f.reg.anio, f.reg.mes)
        reservas[clave].filas.append(f)
    return list(reservas.values())


def _elegir_nombre(filas: list[Fila]) -> NombreLimpio:
    """Prefiere el nombre más confiable y, a igualdad, el de la fuente y fecha más recientes."""
    def clave(f: Fila):
        ref = f.reg.fecha or (date(f.reg.anio, f.reg.mes, 1) if f.reg.anio and f.reg.mes else date.min)
        return (_RANGO_CONFIANZA[f.nom.confianza], ORDEN_FUENTES.index(f.reg.fuente), ref)
    return max(filas, key=clave).nom


def construir_contactos(reservas: list[Reserva]) -> list[Contacto]:
    por_contacto: dict[str, list[Reserva]] = defaultdict(list)
    for r in reservas:
        por_contacto[r.contacto_id].append(r)

    contactos: list[Contacto] = []
    for cid, rs in por_contacto.items():
        todas = [f for r in rs for f in r.filas]
        # Un nombre inválido ("Marzo" en una celda de nombre) no compite: se usa el de otra fila si existe.
        filas = [f for f in todas if f.nom.confianza != INVALIDA] or todas
        invalidas = sorted({f.reg.nombre_raw for f in todas if f.nom.confianza == INVALIDA})
        nombre = _elegir_nombre(filas)
        variantes: list[str] = []
        for f in sorted(filas, key=lambda f: ORDEN_FUENTES.index(f.reg.fuente), reverse=True):
            if f.reg.nombre_raw and f.reg.nombre_raw not in variantes:
                variantes.append(f.reg.nombre_raw)

        # Mismo teléfono con nombres que no comparten ningún token = familia/pareja/tercero o error.
        propios = _tokens(nombre.completo or "")
        conflicto = bool(propios) and any(
            _tokens(f.nom.completo) and not (_tokens(f.nom.completo) & propios) for f in filas)
        confianza = BAJA if conflicto and nombre.confianza in (ALTA, MEDIA) else nombre.confianza

        calidades = {f.tel.calidad for f in filas}
        calidad = VALIDO if VALIDO in calidades else (ASUMIDO_MX if ASUMIDO_MX in calidades else
                                                      (REVISAR if REVISAR in calidades else SIN_TELEFONO))
        notas: list[str] = []
        if nombre.confianza == INVALIDA:
            notas.append(f"Nombre inválido en el origen: {', '.join(repr(v) for v in variantes)}")
        elif invalidas:
            notas.append(f"Una fila traía un nombre inválido ({', '.join(repr(v) for v in invalidas)}); "
                         "se usó el nombre de otra fila con el mismo teléfono")
        if conflicto:
            notas.append("Teléfono compartido por nombres distintos: " + " | ".join(variantes))
        if calidad == ASUMIDO_MX:
            notas.append("Teléfono sin código de país en el origen; se asumió +52")
        if any(f.nom.extras for f in filas):
            notas.append("Anotaciones en el origen: " + " | ".join(sorted({f.nom.extras for f in filas if f.nom.extras})))
        contactos.append(Contacto(cid, None if cid.startswith("NOTEL::") else cid, rs, todas, nombre,
                                  confianza, variantes, conflicto, calidad, notas))
    return contactos


# ------------------------------------------------------------------------------------------------
# Salidas
# ------------------------------------------------------------------------------------------------
def _fechas_ref(c: Contacto) -> list[date]:
    return sorted(r.fecha_ref for r in c.reservas if r.fecha_ref)


def fila_hubspot(c: Contacto) -> dict[str, str]:
    fechas = _fechas_ref(c)
    anios = sorted({str(r.anio_ref) for r in c.reservas if r.anio_ref})
    fuentes = [f for f in ORDEN_FUENTES if any(x.reg.fuente == f for x in c.filas)]
    reservas_ord = sorted(c.reservas, key=lambda r: (r.fecha_ref or date.max, r.etiqueta()))
    n = c.nombre
    sin_nombre = c.confianza == INVALIDA
    return {
        "telefono_e164": c.e164 or "",
        "firstname": "" if sin_nombre else n.nombre,
        "lastname": "" if sin_nombre else n.apellidos,
        "phone": c.e164 or "",
        "lifecyclestage": "customer",
        "nombre_saludo": n.saludo if c.confianza in (ALTA, MEDIA) else "",
        "nombre_confianza": _ETIQUETA_CONFIANZA[c.confianza],
        "nombre_reserva_original": " | ".join(c.variantes),
        "calidad_telefono": c.calidad,
        "whatsapp_opt_in": H.OPT_IN_PENDIENTE,
        "primera_reserva": fechas[0].isoformat() if fechas else "",
        "ultima_reserva": fechas[-1].isoformat() if fechas else "",
        "total_reservas": str(len(c.reservas)),
        "anios_visita": ";".join(anios),
        "fuentes_datos": ";".join(fuentes),
        "resumen_reservas": "; ".join(r.etiqueta() for r in reservas_ord),
        "notas_importacion": " || ".join(c.notas),
    }


def _escribir_csv(path: Path, columnas: list[str], filas: list[dict[str, str]]) -> None:
    # utf-8-sig: Excel muestra bien los acentos; HubSpot ignora el BOM.
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=columnas, lineterminator="\n")
        w.writeheader()
        w.writerows(filas)


def _sugerir_telefono(f: Fila, con_telefono: list[Contacto], por_nombre: dict[str, set[str]]) -> str:
    exactos = por_nombre.get(clave_comparacion(f.nom.completo or f.reg.nombre_raw), set())
    if len(exactos) == 1:
        return f"Coincide por nombre: {next(iter(exactos))}"
    if f.reg.fecha:
        mios = {t for t in _tokens(f.nom.completo or f.reg.nombre_raw) if len(t) >= 4}
        posibles = {c.e164 for c in con_telefono
                    if any(r.fecha == f.reg.fecha for r in c.reservas)
                    and mios & {t for t in _tokens(" ".join(c.variantes)) if len(t) >= 4}}
        if len(posibles) == 1:
            return f"Posible (misma fecha y apellido): {next(iter(posibles))}"
    return ""


def revision(contactos: list[Contacto], filas: list[Fila]) -> dict[str, tuple[list[str], list[list]]]:
    con_tel = [c for c in contactos if c.e164]
    por_nombre: dict[str, set[str]] = defaultdict(set)
    for c in con_tel:
        for v in c.variantes:
            por_nombre[clave_comparacion(limpiar_nombre(v).completo or v)].add(c.e164 or "")

    sin_tel = [f for f in filas if f.tel.calidad == SIN_TELEFONO]
    hojas: dict[str, tuple[list[str], list[list]]] = {}
    hojas["Sin teléfono"] = (
        ["Nombre en el origen", "Fuente", "Ubicación", "Fecha / periodo", "Sugerencia de teléfono"],
        [[f.reg.nombre_raw, f.reg.fuente, f.reg.ubicacion,
          f.reg.fecha.isoformat() if f.reg.fecha else (f.reg.periodo_reporte or f.reg.nota_fecha),
          _sugerir_telefono(f, con_tel, por_nombre)] for f in sin_tel])
    hojas["Teléfono asumido MX"] = (
        ["Nombre", "Teléfono en el origen", "Teléfono E.164", "Fuente(s)"],
        [[c.nombre.completo or c.variantes[0], c.filas[0].reg.telefono_raw, c.e164,
          ", ".join(f for f in ORDEN_FUENTES if any(x.reg.fuente == f for x in c.filas))]
         for c in con_tel if c.calidad == ASUMIDO_MX])
    hojas["Nombres a revisar"] = (
        ["Teléfono", "Nombre(s) en el origen", "Propuesta (nombre / apellidos)", "Confianza", "Motivo"],
        [[c.e164, " | ".join(c.variantes), f"{c.nombre.nombre} / {c.nombre.apellidos}",
          _ETIQUETA_CONFIANZA[c.confianza],
          "Teléfono compartido por nombres distintos" if c.conflicto else
          ("Nombre inválido" if c.confianza == INVALIDA else "No se pudo identificar el nombre de pila")]
         for c in con_tel if c.confianza in (BAJA, INVALIDA)])

    grupos: dict[str, dict[str, list[str]]] = defaultdict(lambda: defaultdict(list))
    for c in con_tel:
        for v in c.variantes:
            grupos[clave_comparacion(limpiar_nombre(v).completo or v)][c.e164 or ""].append(v)
    def _diferencia(tels: list[str]) -> str:
        d = _digitos_distintos(tels[0], tels[1]) if len(tels) == 2 else None
        return ("Difieren en 1 dígito: probable error de captura" if d == 1 else
                f"Difieren en {d} dígitos" if d is not None else "")

    hojas["Mismo nombre, otro teléfono"] = (
        ["Nombre (clave)", "Teléfonos distintos", "Nombres en el origen", "Diferencia"],
        [[k, " , ".join(sorted(tels)), " | ".join(sorted({n for ns in tels.values() for n in ns})),
          _diferencia(sorted(tels))]
         for k, tels in grupos.items() if len(tels) > 1 and k])
    hojas["Reservas sin fecha"] = (
        ["Nombre", "Teléfono", "Fuente", "Ubicación", "Detalle"],
        [[f.reg.nombre_raw, f.tel.e164 or "", f.reg.fuente, f.reg.ubicacion, f.reg.nota_fecha]
         for f in filas if f.reg.precision == S.PRECISION_NINGUNA])
    return hojas


def _escribir_xlsx(path: Path, resumen: list[tuple[str, object]],
                   hojas: dict[str, tuple[list[str], list[list]]]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.append(["Concepto", "Valor"])
    for k, v in resumen:
        ws.append([k, v])
    for nombre, (cab, filas) in hojas.items():
        w = wb.create_sheet(nombre[:31])
        w.append(cab)
        for fila in filas:
            w.append([(" " + v) if isinstance(v, str) and v[:1] in ("=", "@") else v for v in fila])
    for w in wb.worksheets:
        for c in w[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="2F5597")
            c.alignment = Alignment(vertical="center", wrap_text=True)
        w.freeze_panes = "A2"
        w.auto_filter.ref = w.dimensions
        for i, col in enumerate(w.columns, start=1):
            ancho = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            w.column_dimensions[get_column_letter(i)].width = min(max(ancho + 2, 12), 70)
    wb.save(path)


def resumen_estadistico(filas: list[Fila], reservas: list[Reserva], contactos: list[Contacto],
                        auditoria: dict[str, int]) -> list[tuple[str, object]]:
    con_tel = [c for c in contactos if c.e164]
    sin_tel_filas = [f for f in filas if f.tel.calidad == SIN_TELEFONO]
    anios = lambda c: {r.anio_ref for r in c.reservas}  # noqa: E731
    en_2026 = [c for c in con_tel if 2026 in anios(c)]
    previos = [c for c in en_2026 if anios(c) & {2023, 2024, 2025}]
    out: list[tuple[str, object]] = [
        ("Filas leídas (total)", len(filas)),
        *[(f"  · {f}", sum(1 for x in filas if x.reg.fuente == f)) for f in ORDEN_FUENTES],
        ("Reservas distintas (sin duplicados entre hojas/archivos)", len(reservas)),
        ("Filas sin teléfono ('No disponible')", len(sin_tel_filas)),
        ("Contactos con teléfono (a importar)", len(con_tel)),
        ("  · calidad Válido", sum(1 for c in con_tel if c.calidad == VALIDO)),
        ("  · calidad Asumido MX", sum(1 for c in con_tel if c.calidad == ASUMIDO_MX)),
        ("  · calidad Revisar", sum(1 for c in con_tel if c.calidad == REVISAR)),
        ("Contactos con 2 o más reservas (recurrentes)", sum(1 for c in con_tel if len(c.reservas) >= 2)),
        ("Contactos con reserva en 2026", len(en_2026)),
        ("  · de ellos, ya habían estado en 2023-2025", len(previos)),
        ("Nombre confianza Alta", sum(1 for c in con_tel if c.confianza == ALTA)),
        ("Nombre confianza Media", sum(1 for c in con_tel if c.confianza == MEDIA)),
        ("Nombre confianza Baja", sum(1 for c in con_tel if c.confianza == BAJA)),
        ("Nombre inválido", sum(1 for c in con_tel if c.confianza == INVALIDA)),
        ("Teléfonos compartidos por nombres distintos", sum(1 for c in con_tel if c.conflicto)),
        ("Filas corridas en 'Consolidado 2024' (no se usa)", auditoria.get("consolidado_filas_corridas", 0)),
    ]
    por_anio = Counter(a for c in con_tel for a in anios(c) if a)
    out += [(f"Contactos con reservas en {a}", n) for a, n in sorted(por_anio.items())]
    return out


def reporte_markdown(resumen: list[tuple[str, object]]) -> str:
    lineas = ["# Reporte de calidad de datos", "", "| Concepto | Valor |", "|---|---:|"]
    lineas += [f"| {k} | {v} |" for k, v in resumen]
    return "\n".join(lineas) + "\n"


def ejecutar(xlsx_2023_2024: Path, xlsx_2024: Path, pdf_2026: Path, salida: Path) -> list[tuple[str, object]]:
    registros = S.cargar_2023_2024(xlsx_2023_2024)
    r2024, auditoria = S.cargar_2024(xlsx_2024)
    registros += r2024 + S.cargar_2026(pdf_2026)

    filas = normalizar(registros)
    reservas = agrupar_reservas(filas)
    contactos = construir_contactos(reservas)
    importables = sorted((c for c in contactos if c.e164 and c.calidad != REVISAR),
                         key=lambda c: (c.reservas[0].fecha_ref or date.max, c.e164 or ""))

    salida.mkdir(parents=True, exist_ok=True)
    _escribir_csv(salida / "contactos_hubspot.csv", H.COLUMNAS_CSV_CONTACTOS, [fila_hubspot(c) for c in importables])

    historial = []
    for r in sorted(reservas, key=lambda r: (r.contacto_id, r.fecha_ref or date.max)):
        f0 = r.filas[0]
        historial.append({
            "telefono_e164": f0.tel.e164 or "", "nombre": f0.nom.completo or f0.reg.nombre_raw,
            "fecha_estancia": r.fecha.isoformat() if r.fecha else "", "precision_fecha": r.precision,
            "anio": r.anio or "", "mes": r.mes or "",
            "periodos_reporte": ";".join(sorted({x.reg.periodo_reporte for x in r.filas if x.reg.periodo_reporte})),
            "fuentes": ";".join(r.fuentes), "unidades": r.unidades, "filas_origen": len(r.filas),
            "ubicaciones": " | ".join(x.reg.ubicacion for x in r.filas),
        })
    _escribir_csv(salida / "reservas_historial.csv", COLUMNAS_HISTORIAL, historial)

    resumen = resumen_estadistico(filas, reservas, contactos, auditoria)
    _escribir_xlsx(salida / "revision_manual.xlsx", resumen, revision(contactos, filas))
    (salida / "reporte_calidad.md").write_text(reporte_markdown(resumen), encoding="utf-8")
    return resumen


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx-2023-2024", required=True, type=Path)
    ap.add_argument("--xlsx-2024", required=True, type=Path)
    ap.add_argument("--pdf-2026", required=True, type=Path)
    ap.add_argument("--out", default=Path("out"), type=Path)
    args = ap.parse_args(argv)
    for p in (args.xlsx_2023_2024, args.xlsx_2024, args.pdf_2026):
        if not p.is_file():
            ap.error(f"no existe el archivo: {p}")
    for k, v in ejecutar(args.xlsx_2023_2024, args.xlsx_2024, args.pdf_2026, args.out):
        print(f"{k:<62} {v}")
    print(f"\nArchivos generados en {args.out.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
