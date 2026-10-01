"""Prueba de extremo a extremo del generador con archivos sintéticos que reproducen las trampas reales."""
import csv
from datetime import date

import openpyxl
import pytest

from guest_crm import build, hubspot_schema as H, sources as S

reportlab = pytest.importorskip("reportlab")


def _xlsx_2023_2024(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Nombre de la Reserva", "Teléfono", "Día", "Mes", "Año"])
    ws.append(["MARINA QUINTERO LAGOS", "+52 55 1234 5678", 31, "Julio", 2023])
    ws.append(["Pedro Salas", "No disponible", 3, "Agosto", 2023])
    ws.append(["Ana Duarte", "55 2345 6789", 4, "Agosto", 2023])
    ws.append(["Joel Mena", "3345 12:28 42", 4, "Agosto", 2023])      # teléfono con "hora" de Excel
    wb.save(path)


def _xlsx_2024(path):
    wb = openpyxl.Workbook()
    cons = wb.active
    cons.title = "Consolidado 2024"
    cons.append(["Nombre de la Reserva", "Teléfono", "Día", "Mes", "Año", "Mes del Reporte"])
    cons.append(["Ana Duarte", "+52 55 2345 6789", "20", "de Febrero", "2024", "Enero"])
    cons.append(["Rita Soto", "+52 33 3456 7890", "Octubre", "2024", "Octubre", None])  # fila corrida (trampa real)
    for nombre, filas in {
        "Enero 2024": [["Ana Duarte", "+52 55 2345 6789", "20", "de Febrero", "2024"],
                       ["Familia Lagos - Boda Sol", "No disponible", "5", "de Marzo", "2024"]],
        "Febrero 2024": [["Ana Duarte", "+52 55 2345 6789", "20", "de Febrero", "2024"],      # misma reserva reportada otra vez
                         ["Ramírez Soto Gabriela", "+52 81 4567 8901", "No disponible", "No disponible", "No disponible"]],
    }.items():
        ws = wb.create_sheet(nombre)
        ws.append(["Nombre de la Reserva", "Teléfono", "Día", "Mes", "Año"])
        for f in filas:
            ws.append(f)
    ws = wb.create_sheet("Octubre 2024")          # esta hoja no tiene columna Día
    ws.append(["Nombre del Huésped", "Teléfono", "Mes", "Año"])
    ws.append(["Rita Soto", "+52 33 3456 7890", "Octubre", "2024"])
    wb.save(path)


def _pdf_2026(path):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

    datos = [["Nombre", "Teléfono", "Año 2026"],
             ["Marina Quintero", "+52 55 1234\n5678", "Abril"],          # misma persona que la de 2023: recurrente
             ["Marzo", "+52 33 4567\n8912", "Marzo"],                    # celda de nombre con un mes (error real)
             ["Luis Pardo", "+52 664 111\n2222", "Septiembr\ne"]]       # mes partido en dos líneas
    doc = SimpleDocTemplate(str(path), pagesize=letter)
    t = Table(datos, colWidths=[160, 140, 80])
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    doc.build([t])


@pytest.fixture()
def salida(tmp_path):
    a, b, c = tmp_path / "a.xlsx", tmp_path / "b.xlsx", tmp_path / "c.pdf"
    _xlsx_2023_2024(a), _xlsx_2024(b), _pdf_2026(c)
    out = tmp_path / "out"
    resumen = dict(build.ejecutar(a, b, c, out))
    leer = lambda n: list(csv.DictReader((out / n).open(encoding="utf-8-sig")))  # noqa: E731
    return resumen, leer("contactos_hubspot.csv"), leer("reservas_historial.csv"), out


def test_fuentes_se_leen_completas_y_el_consolidado_se_audita(tmp_path):
    b = tmp_path / "b.xlsx"
    _xlsx_2024(b)
    registros, auditoria = S.cargar_2024(b)
    assert len(registros) == 5                                   # solo hojas mensuales, nunca el Consolidado
    assert auditoria["consolidado_filas_corridas"] == 1
    octubre = next(r for r in registros if r.periodo_reporte == "2024-10")
    assert (octubre.precision, octubre.anio, octubre.mes) == (S.PRECISION_MES, 2024, 10)


def test_pdf_une_el_mes_partido_y_conserva_las_filas(tmp_path):
    c = tmp_path / "c.pdf"
    _pdf_2026(c)
    regs = S.cargar_2026(c)
    assert [r.mes for r in regs] == [4, 3, 9]
    assert all(r.precision == S.PRECISION_REPORTE and r.fecha is None for r in regs)


def test_conteos_y_deduplicacion(salida):
    resumen, contactos, reservas, _ = salida
    assert resumen["Filas leídas (total)"] == 4 + 5 + 3
    # Ana Duarte aparece en dos hojas con la misma fecha: una sola reserva (más la de 2023: 2 en total).
    ana = next(c for c in contactos if c["telefono_e164"] == "+525523456789")
    assert ana["total_reservas"] == "2" and ana["primera_reserva"] == "2023-08-04" and ana["ultima_reserva"] == "2024-02-20"
    tels = [c["telefono_e164"] for c in contactos]
    assert len(tels) == len(set(tels))                           # clave única garantizada
    assert all(t.startswith("+") for t in tels)


def test_contacto_recurrente_une_2023_y_2026_por_telefono(salida):
    _, contactos, _, _ = salida
    m = next(c for c in contactos if c["telefono_e164"] == "+525512345678")
    assert m["anios_visita"] == "2023;2026" and m["fuentes_datos"] == "Reservas 2023-2024;Reservas 2026"
    assert m["firstname"] == "Marina" and m["nombre_saludo"] == "Marina"
    assert "2026-04 (mes de reporte)" in m["resumen_reservas"]


def test_nombre_invalido_del_pdf_no_pisa_un_nombre_real_ni_se_importa_como_nombre(salida):
    _, contactos, _, _ = salida
    c = next(c for c in contactos if c["telefono_e164"] == "+523345678912")
    assert c["firstname"] == "" and c["nombre_confianza"] == "Sin nombre" and c["nombre_saludo"] == ""


def test_filas_sin_telefono_no_se_importan_pero_quedan_en_revision(salida):
    resumen, contactos, _, out = salida
    assert resumen["Filas sin teléfono ('No disponible')"] == 2   # Pedro Salas y Familia Lagos
    assert all(c["telefono_e164"] for c in contactos)
    hojas = openpyxl.load_workbook(out / "revision_manual.xlsx")
    assert hojas["Sin teléfono"].max_row - 1 == 2


def test_reserva_sin_fecha_y_telefono_asumido(salida):
    _, contactos, _, _ = salida
    g = next(c for c in contactos if c["telefono_e164"] == "+528145678901")
    assert g["total_reservas"] == "1" and g["primera_reserva"] == "" and "sin fecha" in g["resumen_reservas"]
    assert g["firstname"] == "Gabriela" and g["nombre_confianza"] == "Media"      # "Apellido Apellido Nombre"
    joel = next(c for c in contactos if c["telefono_e164"] == "+523345122842")
    assert joel["calidad_telefono"] == "Asumido MX"


def test_csv_cumple_el_esquema_de_hubspot(salida):
    _, contactos, _, out = salida
    assert list(contactos[0].keys()) == H.COLUMNAS_CSV_CONTACTOS
    for c in contactos:
        assert c["whatsapp_opt_in"] == H.OPT_IN_PENDIENTE        # nadie queda con consentimiento por defecto
        assert c["nombre_confianza"] in H.CONFIANZA_NOMBRE and c["calidad_telefono"] in H.CALIDAD_TELEFONO
        assert set(filter(None, c["anios_visita"].split(";"))) <= set(H.ANIOS)
        assert set(filter(None, c["fuentes_datos"].split(";"))) <= set(H.FUENTES)
        for campo in ("primera_reserva", "ultima_reserva"):
            if c[campo]:
                date.fromisoformat(c[campo])


def test_historial_una_fila_por_reserva_distinta(salida):
    _, _, reservas, _ = salida
    ana = [r for r in reservas if r["telefono_e164"] == "+525523456789"]
    assert len(ana) == 2                                          # 2023-08-04 y 2024-02-20 (reportada en 2 hojas)
    febrero = next(r for r in ana if r["fecha_estancia"] == "2024-02-20")
    assert febrero["filas_origen"] == "2" and febrero["periodos_reporte"] == "2024-01;2024-02"


def test_esquema_hubspot_es_valido():
    nombres = [p["name"] for p in H.PROPIEDADES_CONTACTO]
    assert len(nombres) == len(set(nombres))
    for p in H.PROPIEDADES_CONTACTO:
        assert p["name"].isascii() and p["name"] == p["name"].lower() and len(p["name"]) <= 100
        assert p["name"][0].isalpha() and all(ch.isalnum() or ch == "_" for ch in p["name"])
        for o in p.get("options", []):
            assert o["value"] == o["label"]                      # decisión de diseño documentada
        if p["type"] == "enumeration":
            assert p.get("options")
    assert sum(1 for p in H.PROPIEDADES_CONTACTO if p.get("hasUniqueValue")) == 1     # límite de HubSpot: 10
    assert H.PROPIEDADES_POR_NOMBRE["telefono_e164"]["hasUniqueValue"]
    personalizadas = set(H.COLUMNAS_CSV_CONTACTOS) - set(H.ESTANDAR)
    assert personalizadas <= set(nombres)                        # todo lo que se importa existe como propiedad


def test_pdf_sin_tablas_falla_en_voz_alta_en_lugar_de_importar_cero_filas(tmp_path):
    from reportlab.pdfgen import canvas
    pdf = tmp_path / "vacio.pdf"
    c = canvas.Canvas(str(pdf))
    c.drawString(72, 720, "texto suelto, sin tabla")
    c.save()
    with pytest.raises(ValueError, match="no se detectó ninguna tabla"):
        S.cargar_2026(pdf)


def test_excel_sin_las_columnas_esperadas_falla_con_mensaje_claro(tmp_path):
    ruta = tmp_path / "malo.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["Cliente", "Celular"])
    wb.save(ruta)
    with pytest.raises(ValueError, match="faltan las columnas"):
        S.cargar_2023_2024(ruta)
